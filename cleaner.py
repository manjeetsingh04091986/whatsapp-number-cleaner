"""
Phone number cleaning module for DMG WhatsApp campaign data.
Uses Google's phonenumbers library for per-country validation.
Handles two cases:
  - Files with a Country column (uses that as a hint)
  - Files with just Mobile numbers in international format (auto-detects country)
"""

import re
import pandas as pd
import phonenumbers
from pathlib import Path

# ---------- Excel writer (with text-typed columns) ----------

def _write_excel_text_safe(df, path, text_cols=None):
    """
    Write a DataFrame to xlsx with specified columns forced to TEXT format.
    Prevents Excel from auto-converting long numbers to scientific notation.
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    text_cols = text_cols or []
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    columns = list(df.columns)
    for col_idx, col_name in enumerate(columns, start=1):
        ws.cell(row=1, column=col_idx, value=col_name)

    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        for col_idx, (col_name, value) in enumerate(zip(columns, row), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if value is not None else "")
            if col_name in text_cols:
                cell.number_format = "@"

    for col_idx, col_name in enumerate(columns, start=1):
        if col_name in text_cols:
            col_letter = get_column_letter(col_idx)
            for cell in ws[col_letter]:
                cell.number_format = "@"

    wb.save(path)


COUNTRY_MAP = {
    "uae": "AE", "united arab emirates": "AE", "u.a.e": "AE", "u.a.e.": "AE",
    "saudi arabia": "SA", "ksa": "SA", "saudi": "SA",
    "qatar": "QA", "kuwait": "KW", "bahrain": "BH", "oman": "OM",
    "jordan": "JO", "lebanon": "LB", "iraq": "IQ", "syria": "SY",
    "palestine": "PS", "yemen": "YE",
    "egypt": "EG", "morocco": "MA", "algeria": "DZ", "tunisia": "TN",
    "libya": "LY", "sudan": "SD",
    "india": "IN", "pakistan": "PK", "bangladesh": "BD", "sri lanka": "LK",
    "nepal": "NP",
    "kenya": "KE", "nigeria": "NG", "south africa": "ZA", "ghana": "GH",
    "ethiopia": "ET", "tanzania": "TZ", "uganda": "UG",
    "united kingdom": "GB", "uk": "GB", "u.k.": "GB", "great britain": "GB",
    "united states": "US", "usa": "US", "u.s.a.": "US", "us": "US",
    "germany": "DE", "france": "FR", "italy": "IT", "spain": "ES",
    "china": "CN", "japan": "JP",
    "singapore": "SG", "malaysia": "MY", "indonesia": "ID",
    "philippines": "PH", "thailand": "TH", "australia": "AU",
}


def country_to_iso(country_str):
    if not country_str or pd.isna(country_str):
        return None
    return COUNTRY_MAP.get(str(country_str).strip().lower())


def is_invalid_pattern(digits):
    if not digits:
        return True
    if len(set(digits)) == 1:
        return True
    if len(digits) >= 4 and digits[3:] == "0" * (len(digits) - 3):
        return True
    return False


def validate_and_clean(raw_number, country_iso):
    if not raw_number or pd.isna(raw_number):
        return None
    raw_str = str(raw_number).strip()
    if not raw_str:
        return None

    parsed = None

    if country_iso:
        try:
            parsed = phonenumbers.parse(raw_str, country_iso)
            if not phonenumbers.is_valid_number(parsed):
                parsed = None
        except phonenumbers.NumberParseException:
            parsed = None

    if parsed is None:
        candidate = raw_str if raw_str.startswith("+") else "+" + re.sub(r"\D", "", raw_str)
        try:
            parsed = phonenumbers.parse(candidate, None)
            if not phonenumbers.is_valid_number(parsed):
                return None
        except phonenumbers.NumberParseException:
            return None

    e164 = phonenumbers.format_number(parsed, phonenumbers.PhoneNumberFormat.E164)
    digits = e164.lstrip("+")
    if is_invalid_pattern(digits):
        return None
    return digits


def pick_best_number(mobile_raw, phone_raw, country_iso):
    mobile_cleaned = validate_and_clean(mobile_raw, country_iso)
    if mobile_cleaned:
        return mobile_cleaned, "mobile"
    phone_cleaned = validate_and_clean(phone_raw, country_iso)
    if phone_cleaned:
        return phone_cleaned, "phone"
    return None, "none"


def clean_dataframe(df, mobile_col="Mobile", phone_col="Phone",
                    country_col="Country", keep_cols=None):
    if keep_cols is None:
        keep_cols = ["First name", "Country", "JT"]
    keep_cols_present = [col for col in keep_cols if col in df.columns]

    for col in [mobile_col, phone_col, country_col] + keep_cols:
        if col not in df.columns:
            df[col] = ""

    cleaned_rows = []
    rejected_rows = []
    unknown_countries = set()

    for _, row in df.iterrows():
        mobile_raw = row.get(mobile_col, "")
        phone_raw = row.get(phone_col, "")
        country_raw = row.get(country_col, "")
        country_iso = country_to_iso(country_raw)
        if country_raw and not country_iso:
            unknown_countries.add(str(country_raw).strip())

        if (pd.isna(mobile_raw) or str(mobile_raw).strip() == "") and \
           (pd.isna(phone_raw) or str(phone_raw).strip() == ""):
            rejected_rows.append({**row.to_dict(),
                                  "Rejection_Reason": "Blank/null phone and mobile"})
            continue

        cleaned_number, source = pick_best_number(mobile_raw, phone_raw, country_iso)

        if not cleaned_number:
            mob_digits = re.sub(r"\D", "", str(mobile_raw)) if mobile_raw else ""
            ph_digits = re.sub(r"\D", "", str(phone_raw)) if phone_raw else ""
            if not mob_digits and not ph_digits:
                reason = "Non-numeric / empty after cleaning"
            elif is_invalid_pattern(mob_digits) and (not ph_digits or is_invalid_pattern(ph_digits)):
                reason = "Invalid pattern (repeated digits / all zeros)"
            elif country_iso:
                reason = "Not a valid phone number for " + str(country_iso)
            else:
                reason = "Not a valid international phone number (check country code prefix and length)"
            rejected_rows.append({**row.to_dict(), "Rejection_Reason": reason})
            continue

        out_row = {col: row.get(col, "") for col in keep_cols_present}
        out_row["Mobile"] = cleaned_number
        cleaned_rows.append(out_row)

    cleaned_df = pd.DataFrame(cleaned_rows)
    dedup_before = len(cleaned_df)
    if not cleaned_df.empty:
        cleaned_df = cleaned_df.drop_duplicates(subset=["Mobile"], keep="first")
    dedup_removed = dedup_before - len(cleaned_df)

    rejected_df = pd.DataFrame(rejected_rows)

    stats = {
        "input_rows": len(df),
        "cleaned_rows": len(cleaned_df),
        "rejected_rows": len(rejected_df),
        "duplicates_removed": dedup_removed,
    }
    if not rejected_df.empty:
        stats["rejection_breakdown"] = rejected_df["Rejection_Reason"].value_counts().to_dict()
    if unknown_countries:
        stats["unknown_countries"] = sorted(unknown_countries)

    return cleaned_df, rejected_df, stats


def clean_csv_file(input_path, output_dir, ticket_id,
                   mobile_col="Mobile", phone_col="Phone", country_col="Country"):
    df = pd.read_csv(input_path, dtype=str).fillna("")
    cleaned_df, rejected_df, stats = clean_dataframe(
        df, mobile_col=mobile_col, phone_col=phone_col, country_col=country_col
    )

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    from datetime import datetime
    date_str = datetime.now().strftime("%Y%m%d")

    cleaned_path = output_dir / f"cleaned_{ticket_id}_{date_str}.xlsx"
    rejected_path = output_dir / f"rejected_{ticket_id}_{date_str}.xlsx"

    _write_excel_text_safe(cleaned_df, cleaned_path, text_cols=["Mobile"])
    if not rejected_df.empty:
        _write_excel_text_safe(rejected_df, rejected_path, text_cols=["Mobile", "Phone"])

    stats["cleaned_file"] = str(cleaned_path)
    stats["rejected_file"] = str(rejected_path) if not rejected_df.empty else None
    return stats


def format_summary(stats):
    lines = [
        "=" * 50,
        "Phone cleansing run summary",
        "=" * 50,
        "Input rows:         " + str(stats["input_rows"]),
        "Cleaned rows:       " + str(stats["cleaned_rows"]),
        "Rejected rows:      " + str(stats["rejected_rows"]),
        "Duplicates removed: " + str(stats["duplicates_removed"]),
    ]
    if stats.get("rejection_breakdown"):
        lines.append("")
        lines.append("Rejection reasons:")
        for reason, count in stats["rejection_breakdown"].items():
            lines.append("  - " + str(reason) + ": " + str(count))
    if stats.get("unknown_countries"):
        lines.append("")
        lines.append("Unknown countries (add to COUNTRY_MAP):")
        for c in stats["unknown_countries"]:
            lines.append("  - " + str(c))
    lines.append("")
    lines.append("Cleaned file:  " + str(stats.get("cleaned_file", "n/a")))
    if stats.get("rejected_file"):
        lines.append("Rejected log:  " + str(stats["rejected_file"]))
    lines.append("=" * 50)
    return "\n".join(lines)
